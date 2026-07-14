from __future__ import annotations

import csv
import math
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
BOX = Path(
    "/Users/hrushikeshkalakandra/Library/CloudStorage/Box-Box/"
    "CeMPA shared area/_SimPaths/_SimPathsUK/modules/wealth/00 exploratory analysis"
)
HOUSING_DIR = BOX / "housing wealth/continuous_regression"
MORTGAGE_DIR = BOX / "Mortgage debt/continuous_regression"
MORTGAGE_EP_DIR = BOX / "Mortgage debt/entry_persistence"
OUT_XLSX = ROOT / "input/reg_wealth_housing.xlsx"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def fnum(value):
    if value is None or value == "" or value == ".":
        return None
    x = float(value)
    return None if math.isnan(x) else x


def load_vcv(path: Path) -> tuple[list[str], list[list[float]]]:
    rows = read_csv(path)
    names = [r["variable"] for r in rows]
    cols = list(rows[0].keys())[1:] if rows else []
    matrix = [[float(r[c] or 0) for c in cols] for r in rows]
    return names, matrix


def cov(vcv: tuple[list[str], list[list[float]]], row_name: str, col_name: str) -> float:
    names, matrix = vcv
    try:
        i = names.index(row_name)
        j = names.index(col_name)
    except ValueError:
        return 0.0
    return matrix[i][j]


def matrix_from_vcv(
    vcv: tuple[list[str], list[list[float]]],
    names: list[str],
) -> np.ndarray:
    out = np.zeros((len(names), len(names)))
    for i, row in enumerate(names):
        for j, col in enumerate(names):
            out[i, j] = cov(vcv, row, col)
    return out


def clean_lag_name(name: str) -> str:
    return name if name.startswith("P_") else f"P_{name}"


def java_name_for_variable(name: str) -> str:
    if ":" in name:
        name = name.split(":", 1)[1]
    if name.startswith("year") and name.removeprefix("year").isdigit():
        return f"Year{name.removeprefix('year')}"
    if name.endswith(".year"):
        year = name.split(".", 1)[0].removesuffix("b")
        if year.isdigit():
            return f"Year{year}"
    if name.endswith(".age50plus"):
        age_code = name.split(".", 1)[0].removesuffix("b").removesuffix("o")
        age_names = {
            "5": "Age20to24",
            "6": "Age25to29",
            "7": "Age30to34",
            "8": "Age35to39",
            "10": "Age45to49",
            "11": "Age50plus",
        }
        if age_code in age_names:
            return age_names[age_code]
    if name.endswith(".gor2"):
        region_code = name.split(".", 1)[0].removesuffix("b")
        region_names = {
            "2": "UKD",
            "4": "UKE",
            "5": "UKF",
            "6": "UKG",
            "7": "UKH",
            "8": "UKI",
            "9": "UKJ",
            "10": "UKK",
            "11": "UKL",
            "12": "UKM",
        }
        if region_code in region_names:
            return region_names[region_code]
    if name.endswith(".inc_quintile"):
        q = name.split(".", 1)[0].removesuffix("b")
        if q in {"2", "3", "4", "5"}:
            return f"YdsesC5{q}"

    names = {
        "_cons": "Constant",
        "constant": "Constant",
        "asinh_RW": "AsinhNetHousingWealth",
        "asinh_TW": "AsinhNetNonPensionWealth",
        "new_homeowner": "NewHomeowner",
        "asinh_L_mort_income_ratio": "AsinhLagMortgageDebtToAnnualPrivateIncome",
        "emp": "Employed",
        "na": "NumberMembersOver17",
        "nk": "NumberChildrenAll",
        "nk04": "NumberChildren04",
        "grad": "Graduate",
        "single_female": "SingleFemale",
        "inc_quintile2": "YdsesC52",
        "inc_quintile3": "YdsesC53",
        "inc_quintile4": "YdsesC54",
        "inc_quintile5": "YdsesC55",
        "wave7": "WASRound7",
        "wave8": "WASRound8",
        "reg2": "UKD",
        "reg4": "UKE",
        "reg5": "UKF",
        "reg6": "UKG",
        "reg7": "UKH",
        "reg8": "UKI",
        "reg9": "UKJ",
        "reg10": "UKK",
        "reg11": "UKL",
        "reg12": "UKM",
    }
    return names.get(name, name)


def label_for_variable(name: str) -> str:
    if name == "P_uhat_cont":
        return "Lagged residual (two-year)"
    if name.startswith("P_"):
        return "Lagged " + label_for_variable(name[2:])
    if name.startswith("year") and name.removeprefix("year").isdigit():
        return f"Year {name.removeprefix('year')}"
    if name.endswith(".year"):
        year = name.split(".", 1)[0].removesuffix("b")
        if year.isdigit():
            if "b." in name:
                return f"Year: base ({year})"
            return f"Year {year}"

    labels = {
        "_cons": "Constant",
        "constant": "Constant",
        "asinh_RW": "asinh(net housing wealth)",
        "asinh_TW": "asinh(net non-pension wealth)",
        "ln_mort": "log(mortgage debt)",
        "uhat_cont": "Residual",
        "HousingPersistence": "HousingPersistence",
        "MortgagePersistence": "MortgagePersistence",
        "new_homeowner": "New homeowner",
        "asinh_L_mort_income_ratio": "asinh(lagged mortgage debt / lagged annual private income)",
        "emp": "Employed",
        "na": "Number of adults",
        "nk": "Number of children",
        "nk04": "Children aged 0-4",
        "grad": "Graduate",
        "single_female": "Single female",
        "wave6": "WAS round 6",
        "wave7": "WAS round 7",
        "wave8": "WAS round 8",
        "ep_wave7": "Wave 7 (ref: Wave 6)",
        "ep_wave8": "Wave 8 (ref: Wave 6)",
        "6b.was_round": "WAS round base (6)",
        "7.was_round": "WAS round 7",
        "8.was_round": "WAS round 8",
        "1b.inc_quintile": "Income quintile: base (1)",
        "2.inc_quintile": "Income quintile: 2",
        "3.inc_quintile": "Income quintile: 3",
        "4.inc_quintile": "Income quintile: 4",
        "5.inc_quintile": "Income quintile: 5",
        "inc_quintile2": "Income quintile: 2",
        "inc_quintile3": "Income quintile: 3",
        "inc_quintile4": "Income quintile: 4",
        "inc_quintile5": "Income quintile: 5",
    }
    if name in labels:
        return labels[name]

    if name.endswith(".age50plus"):
        age_code = name.split(".", 1)[0].removesuffix("b")
        age_labels = {
            "4": "Age band: 15-19",
            "5": "Age band: 20-24",
            "6": "Age band: 25-29",
            "7": "Age band: 30-34",
            "8": "Age band: 35-39",
            "9": "Age band: base (40-44)",
            "10": "Age band: 45-49",
            "11": "Age band: 50+",
        }
        return age_labels.get(age_code, name)

    if name.endswith(".pct"):
        decile_code = name.split(".", 1)[0].removesuffix("b")
        if decile_code == "1":
            return "Income decile: base (1)"
        return f"Income decile: {decile_code}"

    if name.startswith("age"):
        age_labels = {
            "age4": "Age band: 15-19",
            "age5": "Age band: 20-24",
            "age6": "Age band: 25-29",
            "age7": "Age band: 30-34",
            "age8": "Age band: 35-39",
            "age10": "Age band: 45-49",
            "age11": "Age band: 50+",
        }
        return age_labels.get(name, name)

    if name.endswith(".gor2"):
        region_code = name.split(".", 1)[0].removesuffix("b")
        region_labels = {
            "1": "Region: base (North East)",
            "2": "Region: North West",
            "4": "Region: Yorkshire and Humber",
            "5": "Region: East Midlands",
            "6": "Region: West Midlands",
            "7": "Region: East of England",
            "8": "Region: London",
            "9": "Region: South East",
            "10": "Region: South West",
            "11": "Region: Wales",
            "12": "Region: Scotland",
        }
        return region_labels.get(region_code, name)

    if name.startswith("reg"):
        region_labels = {
            "reg2": "Region: North West",
            "reg4": "Region: Yorkshire and Humber",
            "reg5": "Region: East Midlands",
            "reg6": "Region: West Midlands",
            "reg7": "Region: East of England",
            "reg8": "Region: London",
            "reg9": "Region: South East",
            "reg10": "Region: South West",
            "reg11": "Region: Wales",
            "reg12": "Region: Scotland",
        }
        return region_labels.get(name, name)

    return name


def relabel_parameter_sheet(ws):
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, str):
            ws.cell(row, 1).value = label_for_variable(value)
    for col in range(3, ws.max_column + 1):
        value = ws.cell(1, col).value
        if isinstance(value, str):
            ws.cell(1, col).value = label_for_variable(value)


def clear_sheet(ws):
    ws.delete_rows(1, ws.max_row)
    if ws.max_column > 1:
        ws.delete_cols(1, ws.max_column)


def write_parameter_sheet(ws, rows: list[dict]):
    clear_sheet(ws)
    ws.cell(1, 1, "REGRESSOR")
    ws.cell(1, 2, "COEFFICIENT")
    for j, row in enumerate(rows, start=3):
        ws.cell(1, j, row["name"])

    for i, row in enumerate(rows, start=2):
        ws.cell(i, 1, row["name"])
        ws.cell(i, 2, row["coef"])
        for j, col in enumerate(rows, start=3):
            value = 0.0
            if row.get("explicit_vcv") is not None and row.get("explicit_vcv") is col.get("explicit_vcv"):
                value = row["explicit_vcv"][row["idx"], col["idx"]]
            elif row.get("vcv") is col.get("vcv") and row.get("base") and col.get("base"):
                value = row["factor"] * col["factor"] * cov(row["vcv"], row["base"], col["base"])
            ws.cell(i, j, value)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = None
    ws.column_dimensions["A"].width = min(44, max(22, max(len(r["name"]) for r in rows) + 2))
    ws.column_dimensions["B"].width = 16
    for col in range(3, min(ws.max_column, 60) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13


def static_rows(coef_path: Path, model: str, vcv_path: Path) -> list[dict]:
    vcv = load_vcv(vcv_path)
    rows = []
    for r in read_csv(coef_path):
        if r.get("model") != model:
            continue
        se = fnum(r.get("std_error") or r.get("se"))
        if se in (None, 0):
            continue
        name = r.get("coefficient_name") or r.get("parameter") or r.get("term")
        current = r.get("current_variable") or name
        rows.append(
            {
                "name": java_name_for_variable(current),
                "coef": fnum(r.get("estimate") or r.get("coef")),
                "base": name,
                "factor": 1.0,
                "vcv": vcv,
            }
        )
    return rows


def logit_rows(coef_path: Path, model: str, vcv_path: Path) -> list[dict]:
    vcv = load_vcv(vcv_path)
    rows = []
    for r in read_csv(coef_path):
        if r.get("model") != model:
            continue
        se = fnum(r.get("se"))
        if se in (None, 0):
            continue
        base = r["term"]
        if ":" in base:
            base = base.split(":", 1)[1]
        rows.append(
            {
                "name": java_name_for_variable(base),
                "coef": fnum(r.get("coef")),
                "base": base,
                "factor": 1.0,
                "vcv": vcv,
            }
        )
    return rows


def expanded_persistence_rows(
    coef_path: Path,
    model: str,
    vcv_path: Path,
    rho2: float,
    lagged_dependent_name: str,
) -> list[dict]:
    vcv = load_vcv(vcv_path)
    beta_rows = []
    for r in read_csv(coef_path):
        if r.get("model") != model:
            continue
        se = fnum(r.get("std_error") or r.get("se"))
        if se in (None, 0):
            continue
        beta_rows.append(r)

    rows: list[dict] = []
    for r in beta_rows:
        base = r.get("coefficient_name") or r.get("parameter") or r.get("term")
        current = r.get("current_variable") or base
        label = r.get("term_label") or label_for_variable(current)
        rows.append(
            {
                "name": label,
                "coef": fnum(r.get("estimate") or r.get("coef")),
                "base": base,
                "factor": 1.0,
                "vcv": vcv,
            }
        )

    rows.append(
        {
            "name": label_for_variable(lagged_dependent_name),
            "coef": rho2,
            "base": None,
            "factor": 0.0,
            "vcv": None,
        }
    )

    for r in beta_rows:
        base = r.get("coefficient_name") or r.get("parameter") or r.get("term")
        current = r.get("current_variable") or base
        lagged = r.get("lagged_variable") or f"P_{current}"
        rows.append(
            {
                "name": label_for_variable(clean_lag_name(lagged)),
                "coef": -rho2 * fnum(r.get("estimate") or r.get("coef")),
                "base": base,
                "factor": -rho2,
                "vcv": vcv,
            }
        )
    return rows


def build_covariate(df: pd.DataFrame, name: str, lag: bool = False) -> np.ndarray:
    prefix = "P_" if lag else ""
    if name == "constant":
        return np.ones(len(df))
    direct = prefix + name
    if direct in df.columns:
        return df[direct].to_numpy(dtype=float)
    if "." in name:
        level, variable = name.split(".", 1)
        level = level.removesuffix("b")
        source = f"P_{variable}" if lag else variable
        if variable in {"age50plus", "gor2", "inc_quintile", "was_round", "year"}:
            return (df[source].to_numpy(dtype=float) == int(level)).astype(float)
    if name.startswith("inc_quintile"):
        q = int(name.removeprefix("inc_quintile"))
        source = "P_inc_quintile" if lag else "inc_quintile"
        return (df[source].to_numpy(dtype=float) == q).astype(float)
    if name.startswith("age"):
        a = int(name.removeprefix("age"))
        source = "P_age50plus" if lag else "age50plus"
        return (df[source].to_numpy(dtype=float) == a).astype(float)
    if name.startswith("reg"):
        g = int(name.removeprefix("reg"))
        source = "P_gor2" if lag else "gor2"
        return (df[source].to_numpy(dtype=float) == g).astype(float)
    if name == "Age29Under":
        source = "P_dvage17" if lag else "dvage17"
        return (df[source].to_numpy(dtype=float) <= 6).astype(float)
    if name.startswith("Age") and "to" in name:
        source = "P_dvage17" if lag else "dvage17"
        ages = name.removeprefix("Age").split("to", 1)
        lo = int(ages[0])
        hi = int(ages[1])
        codes = {
            (30, 34): [7],
            (35, 39): [8],
            (40, 44): [9],
            (45, 49): [10],
            (50, 54): [11],
            (55, 59): [12],
            (60, 64): [13],
            (65, 69): [14],
            (40, 49): [9, 10],
            (50, 59): [11, 12],
        }
        return np.isin(df[source].to_numpy(dtype=float), codes[(lo, hi)]).astype(float)
    if name.startswith("wave"):
        w = int(name.removeprefix("wave"))
        source = "P_was_round" if lag else "was_round"
        return (df[source].to_numpy(dtype=float) == w).astype(float)
    if name.startswith("year") and name.removeprefix("year").isdigit():
        year = int(name.removeprefix("year"))
        source = "P_year" if lag else "year"
        return (df[source].to_numpy(dtype=float) == year).astype(float)
    raise KeyError(f"Cannot construct covariate {prefix}{name}")


def housing_joint_vcv(
    coef_path: Path,
    model: str,
    rho_annual: float,
) -> tuple[list[str], np.ndarray]:
    """Cluster-robust VCV for beta and annual rho in the observed-wave equation.

    The observed WAS transition is approximately two years, so annual rho enters
    the fitted value as rho^2. The workbook row stores annual rho because the
    Java simulation advances one year at a time.
    """
    sample_path = BOX / "housing_continuous_observed_analytic_sample.dta"
    df = pd.read_stata(sample_path, convert_categoricals=False)
    df = df.loc[df["persist_sample"] == 1].copy()

    beta_rows = []
    for r in read_csv(coef_path):
        if r.get("model") != model:
            continue
        se = fnum(r.get("std_error") or r.get("se"))
        if se in (None, 0):
            continue
        beta_rows.append(r)

    names = [r.get("coefficient_name") or r.get("parameter") or r.get("term") for r in beta_rows]
    current_vars = [r.get("current_variable") or n for r, n in zip(beta_rows, names)]
    lagged_vars = [r.get("lagged_variable") or f"P_{v}" for r, v in zip(beta_rows, current_vars)]
    beta = np.array([fnum(r.get("estimate") or r.get("coef")) for r in beta_rows], dtype=float)

    x = np.column_stack([build_covariate(df, v) for v in current_vars])
    xl = np.column_stack([
        build_covariate(df, v[2:], lag=True) if v.startswith("P_") else build_covariate(df, v, lag=True)
        for v in lagged_vars
    ])
    y = df["asinh_RW"].to_numpy(dtype=float)
    ylag = df["P_asinh_RW"].to_numpy(dtype=float)
    weights = df["dwt"].to_numpy(dtype=float)
    clusters = df["was_pid"].to_numpy()

    gamma = rho_annual * rho_annual
    fitted = x @ beta + gamma * (ylag - xl @ beta)
    resid = y - fitted
    jac_beta = x - gamma * xl
    jac_rho = (2.0 * rho_annual * (ylag - xl @ beta)).reshape(-1, 1)
    jac = np.column_stack([jac_beta, jac_rho])

    weighted_jac = jac * weights[:, None]
    bread = jac.T @ weighted_jac
    meat = np.zeros_like(bread)
    for cluster in pd.unique(clusters):
        mask = clusters == cluster
        score = (jac[mask] * (weights[mask] * resid[mask])[:, None]).sum(axis=0)
        meat += np.outer(score, score)

    inv_bread = np.linalg.pinv(bread)
    vcv = inv_bread @ meat @ inv_bread
    n = len(df)
    g = len(pd.unique(clusters))
    p = jac.shape[1]
    if g > 1 and n > p:
        vcv *= (g / (g - 1)) * ((n - 1) / (n - p))
    return names + ["rho_annual"], vcv


def lag_if_consecutive(df: pd.DataFrame, name: str, consecutive: pd.Series) -> pd.Series:
    lagged = df.groupby("was_pid", sort=False)[name].shift(1)
    return lagged.where(consecutive)


def mortgage_joint_vcv(
    coef_path: Path,
    model: str,
    rho_annual: float,
) -> tuple[list[str], np.ndarray]:
    raw_path = BOX / "was_wealthdata_5_8_longitudinal.dta"
    keep = [
        "was_pid", "was_round", "year", "bu_rp", "dwt", "dvage17", "emp", "na",
        "nk", "nk04", "gor2", "grad", "pct", "sex", "wealth", "tot_open",
        "tot_pp", "dvhvalue", "main_mort",
    ]
    df = pd.read_stata(raw_path, columns=keep, convert_categoricals=False)
    df = df.loc[df["bu_rp"] == 1].copy()
    required = [
        "was_pid", "was_round", "year", "dwt", "dvage17", "emp", "na", "nk", "nk04",
        "gor2", "grad", "pct", "sex", "wealth", "tot_open", "tot_pp",
        "dvhvalue", "main_mort",
    ]
    df = df.dropna(subset=required)
    df = df.loc[df["dwt"] > 0].copy()
    df["single_female"] = ((df["na"] == 1) & (df["sex"] == 2)).astype(float)
    df["age50plus"] = df["dvage17"].clip(lower=4, upper=11)
    df["inc_quintile"] = np.ceil(df["pct"] / 2)
    df["TW"] = df["wealth"] - df["tot_open"] - df["tot_pp"]
    df["RW"] = df["dvhvalue"] - df["main_mort"]
    df["homeowner_gross"] = (df["dvhvalue"] > 0).astype(float)
    df["has_mortgage"] = ((df["main_mort"] > 0) & (df["homeowner_gross"] == 1)).astype(float)
    df["ln_mort"] = np.nan
    mort_mask = df["has_mortgage"] == 1
    df.loc[mort_mask, "ln_mort"] = np.log(df.loc[mort_mask, "main_mort"])
    df["asinh_RW"] = np.arcsinh(df["RW"])
    df["asinh_TW"] = np.arcsinh(df["TW"])

    df = df.sort_values(["was_pid", "was_round"]).copy()
    prev_round = df.groupby("was_pid", sort=False)["was_round"].shift(1)
    consecutive = df["was_round"].eq(prev_round + 1)
    for name in [
        "was_round", "year", "RW", "TW", "homeowner_gross", "has_mortgage", "dvage17", "age50plus",
        "emp", "na", "nk", "nk04", "gor2", "grad", "single_female",
        "inc_quintile", "ln_mort", "asinh_RW", "asinh_TW",
    ]:
        df[f"P_{name}"] = lag_if_consecutive(df, name, consecutive)

    df["new_mortgage"] = np.where(
        df[["has_mortgage", "P_has_mortgage"]].notna().all(axis=1),
        ((df["has_mortgage"] == 1) & (df["P_has_mortgage"] == 0)).astype(float),
        np.nan,
    )
    df["continuing_mortgage"] = np.where(
        df[["has_mortgage", "P_has_mortgage"]].notna().all(axis=1),
        ((df["has_mortgage"] == 1) & (df["P_has_mortgage"] == 1)).astype(float),
        np.nan,
    )
    target_required = [
        "ln_mort", "asinh_RW", "asinh_TW", "new_mortgage", "age50plus", "emp",
        "na", "nk", "nk04", "gor2", "grad", "single_female", "inc_quintile",
        "year", "dwt",
    ]
    df["target_sample"] = (df["has_mortgage"] == 1) & df[target_required].notna().all(axis=1)
    df["cont_target_sample"] = df["target_sample"] & (df["continuing_mortgage"] == 1)
    df = df.loc[df["cont_target_sample"] & df["P_ln_mort"].notna()].copy()

    beta_rows = []
    for r in read_csv(coef_path):
        if r.get("model") != model:
            continue
        se = fnum(r.get("std_error") or r.get("se"))
        if se in (None, 0):
            continue
        beta_rows.append(r)

    names = [r.get("coefficient_name") or r.get("parameter") or r.get("term") for r in beta_rows]
    current_vars = [r.get("current_variable") or n for r, n in zip(beta_rows, names)]
    beta = np.array([fnum(r.get("estimate") or r.get("coef")) for r in beta_rows], dtype=float)

    x = np.column_stack([build_covariate(df, v) for v in current_vars])
    xl = np.column_stack([build_covariate(df, v, lag=True) for v in current_vars])
    y = df["ln_mort"].to_numpy(dtype=float)
    ylag = df["P_ln_mort"].to_numpy(dtype=float)
    weights = df["dwt"].to_numpy(dtype=float)
    clusters = df["was_pid"].to_numpy()

    gamma = rho_annual * rho_annual
    fitted = x @ beta + gamma * (ylag - xl @ beta)
    resid = y - fitted
    jac_beta = x - gamma * xl
    jac_rho = (2.0 * rho_annual * (ylag - xl @ beta)).reshape(-1, 1)
    jac = np.column_stack([jac_beta, jac_rho])

    weighted_jac = jac * weights[:, None]
    bread = jac.T @ weighted_jac
    meat = np.zeros_like(bread)
    for cluster in pd.unique(clusters):
        mask = clusters == cluster
        score = (jac[mask] * (weights[mask] * resid[mask])[:, None]).sum(axis=0)
        meat += np.outer(score, score)

    inv_bread = np.linalg.pinv(bread)
    vcv = inv_bread @ meat @ inv_bread
    n = len(df)
    g = len(pd.unique(clusters))
    p = jac.shape[1]
    if g > 1 and n > p:
        vcv *= (g / (g - 1)) * ((n - 1) / (n - p))
    return names + ["rho_annual"], vcv


def expanded_persistence_rows_delta(
    coef_path: Path,
    model: str,
    gamma: float,
    lagged_dependent_name: str,
    beta_vcv_path: Path | None = None,
    gamma_var: float | None = None,
    joint_vcv: tuple[list[str], np.ndarray] | None = None,
) -> list[dict]:
    beta_rows = []
    for r in read_csv(coef_path):
        if r.get("model") != model:
            continue
        se = fnum(r.get("std_error") or r.get("se"))
        if se in (None, 0):
            continue
        beta_rows.append(r)

    base_names = [r.get("coefficient_name") or r.get("parameter") or r.get("term") for r in beta_rows]
    beta = np.array([fnum(r.get("estimate") or r.get("coef")) for r in beta_rows], dtype=float)
    k = len(beta)

    if joint_vcv is not None:
        joint_names, param_vcv = joint_vcv
        idx = [joint_names.index(name) for name in base_names] + [joint_names.index("gamma")]
        param_vcv = param_vcv[np.ix_(idx, idx)]
    else:
        if beta_vcv_path is None or gamma_var is None:
            raise ValueError("beta_vcv_path and gamma_var are required without joint_vcv")
        beta_vcv = matrix_from_vcv(load_vcv(beta_vcv_path), base_names)
        param_vcv = np.zeros((k + 1, k + 1))
        param_vcv[:k, :k] = beta_vcv
        param_vcv[k, k] = gamma_var

    transform = np.zeros((2 * k + 1, k + 1))
    transform[:k, :k] = np.eye(k)
    transform[k, k] = 1.0
    for i in range(k):
        transform[k + 1 + i, i] = -gamma
        transform[k + 1 + i, k] = -beta[i]
    expanded_vcv = transform @ param_vcv @ transform.T

    rows: list[dict] = []
    for i, r in enumerate(beta_rows):
        current = r.get("current_variable") or base_names[i]
        rows.append(
            {
                "name": r.get("term_label") or label_for_variable(current),
                "coef": beta[i],
                "idx": i,
                "explicit_vcv": expanded_vcv,
            }
        )

    rows.append(
        {
            "name": label_for_variable(lagged_dependent_name),
            "coef": gamma,
            "idx": k,
            "explicit_vcv": expanded_vcv,
        }
    )

    for i, r in enumerate(beta_rows):
        current = r.get("current_variable") or base_names[i]
        lagged = r.get("lagged_variable") or f"P_{current}"
        rows.append(
            {
                "name": label_for_variable(clean_lag_name(lagged)),
                "coef": -gamma * beta[i],
                "idx": k + 1 + i,
                "explicit_vcv": expanded_vcv,
            }
        )
    return rows


def h1_persistence_rows_delta(
    coef_path: Path,
    model: str,
    rho_annual: float,
    persistence_name: str,
    beta_vcv_path: Path | None = None,
    rho_var: float | None = None,
    joint_vcv: tuple[list[str], np.ndarray] | None = None,
) -> list[dict]:
    """Return current-period beta rows plus one annual residual-persistence row."""
    beta_rows = []
    for r in read_csv(coef_path):
        if r.get("model") != model:
            continue
        se = fnum(r.get("std_error") or r.get("se"))
        if se in (None, 0):
            continue
        beta_rows.append(r)

    base_names = [r.get("coefficient_name") or r.get("parameter") or r.get("term") for r in beta_rows]
    beta = np.array([fnum(r.get("estimate") or r.get("coef")) for r in beta_rows], dtype=float)
    k = len(beta)

    if joint_vcv is not None:
        joint_names, param_vcv = joint_vcv
        idx = [joint_names.index(name) for name in base_names] + [joint_names.index("rho_annual")]
        param_vcv = param_vcv[np.ix_(idx, idx)]
    else:
        if beta_vcv_path is None or rho_var is None:
            raise ValueError("beta_vcv_path and rho_var are required without joint_vcv")
        beta_vcv = matrix_from_vcv(load_vcv(beta_vcv_path), base_names)
        param_vcv = np.zeros((k + 1, k + 1))
        param_vcv[:k, :k] = beta_vcv
        param_vcv[k, k] = rho_var

    rows: list[dict] = []
    for i, r in enumerate(beta_rows):
        current = r.get("current_variable") or base_names[i]
        rows.append(
            {
                "name": java_name_for_variable(current),
                "coef": beta[i],
                "idx": i,
                "explicit_vcv": param_vcv,
            }
        )

    rows.append(
        {
            "name": label_for_variable(persistence_name),
            "coef": rho_annual,
            "idx": k,
            "explicit_vcv": param_vcv,
        }
    )
    return rows


def get_summary_value(path: Path, model: str, field: str) -> float:
    for row in read_csv(path):
        if row.get("model") == model:
            value = fnum(row.get(field))
            if value is None:
                raise ValueError(f"Missing {field} for {model} in {path}")
            return value
    raise ValueError(f"Missing model {model} in {path}")


def update_info(ws):
    ws["B3"] = "09/06/2026 HK/Codex (H1-style residual-persistence packaging; mortgage value uses stable two-step persistence)"
    labels = {
        "HW2a": "Mortgage holding: entry logit",
        "HW2b": "Mortgage holding: persistence logit",
        "HW1c": "Net housing value: continuing homeowners, H1-style residual-persistence equation",
        "HW1d": "Net housing value: new homeowners, static equation",
        "HW2c": "Mortgage value: continuing mortgage holders, H1-style residual-persistence equation",
        "HW2d": "Mortgage value: new mortgage holders, static equation",
    }
    existing = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    for key, label in labels.items():
        row = existing.get(key)
        if row:
            ws.cell(row, 2, label)


def update_gof_continuous(ws, housing_summary_path: Path, mortgage_summary_path: Path):
    for r in range(ws.max_row, 0, -1):
        label = ws.cell(r, 1).value
        if isinstance(label, str) and ("xsect_probit" in label.lower() or "probit" in label.lower()):
            ws.delete_rows(r, 3)

    rows = {row["model"]: row for row in read_csv(housing_summary_path)}
    rows.update({row["model"]: row for row in read_csv(mortgage_summary_path)})
    mapping = {
        "HW1c - Net housing value: continuing homeowners": "hcw_persist",
        "HW1d - Net housing value: new homeowners": "hcw_new",
        "HW2c - Mortgage value: continuing mortgage holders": "continuing_mortgage_target",
        "HW2d - Mortgage value: new mortgage holders": "new_mortgage_target",
    }

    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label not in mapping:
            continue
        source = rows[mapping[label]]
        for row_idx, values in {
            r: [label, "N", fnum(source.get("n")), "R-squared / pseudo R-squared", fnum(source.get("r2"))],
            r + 1: [None, "RMSE", fnum(source.get("rmse")), "rho annual", fnum(source.get("rho_annual"))],
            r + 2: [None, "sigma2 eta", fnum(source.get("sigma2_eta")), "sigma2 epsilon annual", fnum(source.get("sigma2_eps_annual"))],
        }.items():
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row_idx, col_idx).value = value


def update_auxiliary(ws, housing_summary_path: Path, mortgage_summary_path: Path):
    housing = next(row for row in read_csv(housing_summary_path) if row["model"] == "hcw_persist")
    mortgage = next(
        row for row in read_csv(mortgage_summary_path)
        if row["model"] == "continuing_mortgage_target"
    )
    updates = {
        ("HW1c", "persist_rho_annual"): fnum(housing.get("rho_annual")),
        ("HW1c", "persist_rho_sq_annual_link"): fnum(housing.get("rho_sq_annual_link")),
        ("HW1c", "persist_sigma2_eta"): fnum(housing.get("sigma2_eta")),
        ("HW1c", "persist_sigma2_eps_annual"): fnum(housing.get("sigma2_eps_annual")),
        ("HW2c", "cont_rho_2yr"): fnum(mortgage.get("rho_2yr")),
        ("HW2c", "cont_rho_annual"): fnum(mortgage.get("rho_annual")),
        ("HW2c", "cont_sigma2_eta"): fnum(mortgage.get("sigma2_eta")),
        ("HW2c", "cont_sigma2_eps_annual"): fnum(mortgage.get("sigma2_eps_annual")),
        ("HW2c", "cont_annualisation_stable"): fnum(mortgage.get("annualisation_stable")),
    }
    for r in range(2, ws.max_row + 1):
        key = (ws.cell(r, 1).value, ws.cell(r, 2).value)
        if key in updates:
            ws.cell(r, 3, updates[key])
        if key == ("HW2c", "cont_annualisation_stable"):
            ws.cell(r, 4).value = (
                "Equals 1 if the recovered annual residual-persistence parameter "
                "is defined and below one; equals 0 otherwise."
            )


def update_gof_mortgage(ws, mortgage_summary_path: Path):
    rows = {row["model"]: row for row in read_csv(mortgage_summary_path)}
    mapping = {
        "HW2c - Mortgage value: continuing mortgage holders": "continuing_mortgage_target",
        "HW2d - Mortgage value: new mortgage holders": "new_mortgage_target",
    }

    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label not in mapping:
            continue
        source = rows[mapping[label]]
        for row_idx, values in {
            r: [label, "N", fnum(source.get("n")), "R-squared / pseudo R-squared", fnum(source.get("r2"))],
            r + 1: [None, "RMSE", fnum(source.get("rmse")), "rho annual", fnum(source.get("rho_annual"))],
            r + 2: [None, "sigma2 eta", fnum(source.get("sigma2_eta")), "sigma2 epsilon annual", fnum(source.get("sigma2_eps_annual"))],
        }.items():
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row_idx, col_idx).value = value


def update_gof_mortgage_incidence(ws, mortgage_ep_summary_path: Path):
    rows = {row["model"]: row for row in read_csv(mortgage_ep_summary_path)}
    mapping = {
        "HW2a - Prob. Mortgage entry": "mort_entry_logit",
        "HW2b - Prob. Mortgage persistence": "mort_persist_logit",
    }

    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label not in mapping:
            continue
        source = rows[mapping[label]]
        for row_idx, values in {
            r: [
                label,
                "N",
                fnum(source.get("n")),
                "R-squared / pseudo R-squared",
                fnum(source.get("r2")),
            ],
            r + 1: [
                None,
                "Log likelihood",
                fnum(source.get("log_likelihood")),
                "Chi^2 / F",
                fnum(source.get("chi2_or_F")),
            ],
        }.items():
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row_idx, col_idx).value = value


def update_auxiliary_mortgage(ws, mortgage_summary_path: Path):
    mortgage = next(
        row for row in read_csv(mortgage_summary_path)
        if row["model"] == "continuing_mortgage_target"
    )
    updates = {
        ("HW2c", "cont_rho_2yr"): fnum(mortgage.get("rho_2yr")),
        ("HW2c", "cont_rho_annual"): fnum(mortgage.get("rho_annual")),
        ("HW2c", "cont_sigma2_eta"): fnum(mortgage.get("sigma2_eta")),
        ("HW2c", "cont_sigma2_eps_annual"): fnum(mortgage.get("sigma2_eps_annual")),
        ("HW2c", "cont_annualisation_stable"): fnum(mortgage.get("annualisation_stable")),
    }
    for r in range(2, ws.max_row + 1):
        key = (ws.cell(r, 1).value, ws.cell(r, 2).value)
        if key in updates:
            ws.cell(r, 3, updates[key])
        if key == ("HW2c", "cont_annualisation_stable"):
            ws.cell(r, 4).value = (
                "Equals 1 if the recovered annual residual-persistence parameter "
                "is defined and below one; equals 0 otherwise."
            )


def main():
    wb = openpyxl.load_workbook(OUT_XLSX)

    mortgage_ep_coef = MORTGAGE_EP_DIR / "mortgage_ep_model_results.csv"
    write_parameter_sheet(
        wb["HW2a"],
        logit_rows(
            mortgage_ep_coef,
            "mort_entry_logit",
            MORTGAGE_EP_DIR / "mortgage_ep_vcv_mort_entry_logit.csv",
        ),
    )
    write_parameter_sheet(
        wb["HW2b"],
        logit_rows(
            mortgage_ep_coef,
            "mort_persist_logit",
            MORTGAGE_EP_DIR / "mortgage_ep_vcv_mort_persist_logit.csv",
        ),
    )

    mortgage_coef = MORTGAGE_DIR / "mortgage_continuous_coefficients.csv"
    mortgage_summary = MORTGAGE_DIR / "mortgage_continuous_model_summary.csv"
    mortgage_rho_annual = get_summary_value(mortgage_summary, "continuing_mortgage_target", "rho_annual")
    write_parameter_sheet(
        wb["HW2c"],
        h1_persistence_rows_delta(
            mortgage_coef,
            "mort_cont_target",
            mortgage_rho_annual,
            "MortgagePersistence",
            joint_vcv=mortgage_joint_vcv(mortgage_coef, "mort_cont_target", mortgage_rho_annual),
        ),
    )
    write_parameter_sheet(
        wb["HW2d"],
        static_rows(
            mortgage_coef,
            "mort_new_target",
            MORTGAGE_DIR / "mortgage_continuous_vcv_mort_new_target.csv",
        ),
    )

    for sheet in ["HW2a", "HW2b", "HW2c", "HW2d"]:
        if sheet in wb.sheetnames:
            relabel_parameter_sheet(wb[sheet])

    update_info(wb["Info"])
    update_gof_mortgage_incidence(
        wb["Gof"],
        MORTGAGE_EP_DIR / "mortgage_ep_model_summary.csv",
    )
    update_gof_mortgage(wb["Gof"], mortgage_summary)
    update_auxiliary_mortgage(wb["Auxiliary parameters"], mortgage_summary)
    wb.save(OUT_XLSX)


if __name__ == "__main__":
    main()
